import re
import os
from collections import Counter
from docx import Document

from utils import sanitize_string, REQUIRED_COLUMNS
from locales import STRINGS

class Result:
    def __init__(self, fields):
        self.ok = True
        self.message = None

        self.out_fname = ''
        self.fields = fields
        self.match_count = Counter()

        self.name = (fields if fields is not None else {}).get('name', None)

    def set_error(self, message):
        self.ok = False
        self.message = message
        return self

    def add_matches_count(self, text, placeholders):
        for placeholder in placeholders:
            if placeholder in text:
                self.match_count[placeholder] += 1

    def __str__(self):
        if self.ok:
            return f'{self.fields["name"]} counter={sum(self.match_count.values())} out={os.path.basename(self.out_fname)}'
        return f'error: {self.message}'

    def __repr__(self):
        return self.__str__()


_FIELDS_PLACEHOLDERS = {
    'name': 'NOMBRECOMPLETO',
    'rut': 'RUT',
    'home_address': 'DOMICILIO',
    'date_gen': 'FECHAEMISION',
    'date_start': 'FECHAINICIO',
    'date_end': 'FECHATERMINO',
    'amount_num': 'MONTONUMERO',
    'amount_words': 'MONTOPALABRAS',
}
_PLACEHOLDER_REGEX = re.compile('|'.join(_FIELDS_PLACEHOLDERS.values()))

assert set(_FIELDS_PLACEHOLDERS.keys()).issubset(REQUIRED_COLUMNS), "Internal error: field in _FIELDS_PLACEHOLDER is not marked as required"

def fill_document(fields, out_folder='generated', template_folder='templates'):
    result = Result(fields)

    if "template" not in fields:
        return result.set_error(message=STRINGS.get('missing_col', 'template'))

    template_fname = os.path.join(template_folder, f'{fields["template"]}.docx')

    if not os.path.isfile(template_fname):
        return result.set_error(message=STRINGS.get('template_file_not_found', template_fname))

    doc = Document(template_fname)

    if any(key not in fields for key in _FIELDS_PLACEHOLDERS.keys()):
        missing_fields = ','.join(key for key in _FIELDS_PLACEHOLDERS.keys() if key not in fields)
        return result.set_error(message=STRINGS.get("missing_cols", missing_fields))

    for p in doc.paragraphs:
        for r in p.runs:
            if not re.search(_PLACEHOLDER_REGEX, r.text):
                continue

            result.add_matches_count(r.text, _FIELDS_PLACEHOLDERS.values())

            replaced = str(r.text)
            for field, placeholder in _FIELDS_PLACEHOLDERS.items():
                replaced = replaced.replace(placeholder, fields[field])

            r.text = replaced

    out_fname = os.path.join(out_folder, f'{sanitize_string(fields["name"])}.docx')
    os.makedirs(os.path.dirname(out_fname), exist_ok=True)
    doc.save(out_fname)

    result.out_fname = out_fname

    return result
